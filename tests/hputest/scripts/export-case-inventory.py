#!/usr/bin/env python3
"""导出真实参数及留空的IT/ST验收表；不把编译/模型通过写成硬件PASS。"""
import argparse
import base64
import csv
import json
import math
from pathlib import Path
import re
import struct
import subprocess

ROOT = Path(__file__).resolve().parents[1]
REPO = ROOT.parents[1]
GROUPS = ["冒烟", "配置、连接与数据通路", "边界与CPU交叉", "基础指令", "基本/组合算子", "SEAL算法库", "完整应用", "N65536性能"]


def table(path, delimiter="\t"):
    with path.open() as stream: return list(csv.DictReader(stream, delimiter=delimiter))


def load(path):
    return json.loads(path.read_text())


def product(values):
    return str(math.prod(values)) if values else "不适用"


def revision(path):
    return subprocess.check_output(["git", "-C", str(path), "rev-parse", "HEAD"], text=True).strip()


def application(spec, generated, native, cache):
    stem=spec["program_stem"]; directory=generated/"application-data"/stem
    upstream=directory/"upstream"; pkg=load(upstream/"package.json")
    params=load(upstream/pkg["parameters"])
    graph=load(upstream/pkg["operation_graph"])["operations"]
    memory=load(upstream/pkg["memory_config"])
    allocations=table(upstream/pkg["line_map"],",")
    q=params["levels"][0]["moduli"]; p=params["key_moduli"][len(q):]
    n=params["poly_modulus_degree"]
    extra={}
    if spec["role"]!="application":
        parameter_file=directory/"FIXTURE_PARAMETERS.json"
        if not parameter_file.exists():
            parameter_file=cache/(stem+".json")
            if not parameter_file.exists():
                profile="library" if stem.startswith("seal_") else "regression" if "/04_parameter_regression/" in spec["source"] else "kernel"
                parameter_file.write_bytes(subprocess.check_output([str(native),spec["scheme"],spec["role"],str(n),profile]))
        extra=load(parameter_file)
        assert extra["degree"]==n and extra["Q_rns_count"]==len(q)
    digits={int(m[1]) for row in allocations if row["kind"]=="evaluation_key"
            for m in re.finditer(r"/(?:d|digit)(\d+)/",row["allocation_id"])}
    if extra: assert len(digits)==extra["dnum"], f"{stem}: key digit metadata differs from native SEAL"
    mod_allocation=next(row for row in allocations if row["kind"]=="modulus_table")
    with (upstream/pkg["image"]).open("rb") as stream:
        stream.seek(int(mod_allocation["line_offset"])*256)
        raw=stream.read(int(mod_allocation["payload_words"])*4)
    words=struct.unpack(f"<{len(raw)//4}I",raw)
    mod_contexts=list(words[::4])
    instructions=(directory/(stem+".inst32")).read_text().split()
    dma=table(directory/"dma_relocation_manifest.csv",",")
    normalized_lines=(directory/"application_window.u32.bin").stat().st_size//256
    input_limbs={}
    for row in allocations:
        if row["kind"]=="ciphertext" and row["read_only"]=="true":
            name=row["allocation_id"].split('/c')[0]
            input_limbs[name]=input_limbs.get(name,0)+1
    recipe=("CKKS slots A=[0.25,-1.5,2,0.75], B=[2,0.5,-1.25,1.5]，其余0" if spec['scheme']=='ckks' else
            "BatchEncoder A[i]=i%5+1, B[i]=i%3+1, i=0..N-1") if extra else "上游完整应用包，见operation_graph及固定生产者example"
    return dict(N=n,scheme=spec["scheme"],operation=spec["role"],Q=q,P=p,
        L=params["levels"][0]["chain_index"],dnum=len(digits) or None,
        output_Q=[len(next(level["moduli"] for level in params["levels"] if level["parms_id"]==op["parms_id"])) for op in graph],
        plain_modulus=params.get("plain_modulus",0),key_moduli=params["key_moduli"],levels=params["levels"],
        security_level_bits=params["security_level_bits"],operation_graph=graph,fixture=extra,input_recipe=recipe,
        input_domain="coefficient/bit-reversed" if spec['scheme']=='bfv' else "NTT/P-network physical",
        rotation_galois_element=pow(3,extra['rotation_steps'],2*n) if extra.get('rotation_steps') else None,
        input_limb_counts=input_limbs,modulus_table=mod_contexts,
        evaluation_key_limbs=sum(r["kind"]=="evaluation_key" for r in allocations),
        degree_source="package metadata/parameters.json",golden_source="modified-SEAL Evaluator, all raw physical words",
        producer_commit=(directory/"producer_commit.txt").read_text().strip(),program_stem=stem,
        memory_base="0x90000000" if stem.endswith("_n65536") else "0x87000000",
        used_lines=normalized_lines-64,guard_lines=64,image_bytes=normalized_lines*256,
        golden_bytes=(directory/"application_golden.u32.bin").stat().st_size,
        instruction_count=len(instructions),dma_count=len(dma),
        producer_used_lines=memory["used_lines"],minimum_ram_bytes=536870912 if stem.endswith("_n65536") else 268435456,
        source_parameters=str((directory.relative_to(generated)/"upstream"/pkg["parameters"]).as_posix()),
        model_verified=True,rtl_verified=False)


def legacy_parameters(case,generated,legacy):
    cid=case["case_id"]
    if cid in ("04_01_bconv_q4_to_p3_n4096","04_02_ntt_n4096","04_03_intt_n4096","04_10_legacy_auto_n4096"):
        op={"04_01":"bconv","04_02":"ntt","04_03":"intt","04_10":"auto"}[cid[:5]]
        path=legacy/op/"test_data/params.json"
        params=load(path); q=params["moduli"]
        return dict(N=params["N"],scheme="legacy RNS",operation=op,Q=q[:4],P=q[4:],L=None,
                    dnum=2 if op=="auto" else None,producer_parameters=params,
                    memory_base="0x87000000",plain_modulus=None,security_level_bits=0,
                    producer_commit=revision(ROOT/"third_party/inline-asm"),golden_source="inline-asm FHE reference",
                    degree_source="legacy test_data/params.json")
    if case["qualifier"] in ("blocked-not-issued","termination-probe-pass","termination-probe-fail"):
        return dict(N=None,Q=[],P=[],L=None,dnum=None,degree_source="无已执行计算；不填猜测参数")
    # 控制/边界用例仅使用一个RNS的A/B夹具，不是SEAL密文。
    text=(ROOT/case["source"]).read_text()
    q=[50061313,50077697] if "hpu/it_v2.h" in text or "hpu/steps.h" in text else [50061313]
    return dict(N=4096,scheme="raw RNS / control",Q=q[:1],P=[],modulus_table=q,L=None,dnum=None,
                operation=cid[6:],memory_base="0x87000000",plain_modulus=None,security_level_bits=0,
                producer_commit=revision(ROOT/"third_party/inline-asm"),golden_source="C modular reference / status invariants",
                degree_source="N4096 fixture；列出的q为候选上下文，不代表多RNS密文",
                note="纯CSR/握手项无密文语义；以源码及subcase检查点为准。")


def collect(args):
    cases=table(ROOT/"cases.tsv"); schemes={r["case_id"]:r for r in table(ROOT/"scheme-cases.tsv")}
    large={r["case_id"]:r for r in table(ROOT/"large-cases.tsv")}
    perf={r["case_id"]:r for r in table(ROOT/"perf-cases.tsv")}
    cache=args.output/"fixture-parameters";cache.mkdir(parents=True,exist_ok=True)
    records=[]; by_id={}
    for row in cases:
        cid=row["case_id"]
        if cid in schemes: params=application(schemes[cid],args.generated,args.parameter_tool,cache)
        elif cid in large:
            directory=args.generated/"large-data"/large[cid]["program_stem"]
            params=load(directory/"PARAMETERS.json")
            params.update(N=params["degree"],Q=params["q_moduli"],P=params["p_moduli"],L=None,
                          producer_commit=params["legacy_commit"],degree_source="N65536 producer + independent model")
            params.update(image_bytes=(directory/"application_window.u32.bin").stat().st_size,
                          golden_bytes=(directory/"application_golden.u32.bin").stat().st_size)
            if params['operation']=='pmodld':
                params.update(Q=params['q_moduli'][:1],modulus_table=params['q_moduli'],
                              note='三个串接子操作选择q0/q1/q0；每次仅1个RNS，不是Q2密文。')
        elif cid in perf:
            params=dict(by_id[perf[cid]["functional_case_id"]]["parameters"])
            params.update(functional_case_id=perf[cid]["functional_case_id"],metric="SCPU cycle: first HPU issue to completion")
        else: params=legacy_parameters(row,args.generated,args.legacy)
        item=dict(row,parameters=params,IT_pass=None,ST_pass=None,measured_seconds=None,
                  tested_AM_revision=None,tested_RTL_revision=None)
        records.append(item);by_id[cid]=item
    return records,by_id


def fetch_st(gh):
    commit=subprocess.check_output([gh,"api","repos/cuihu2/System-test/commits/main","--jq",".sha"],text=True).strip()
    payload=json.loads(subprocess.check_output([gh,"api",f"repos/cuihu2/System-test/contents/CASE_STATUS.md?ref={commit}"],text=True))
    text=base64.b64decode(payload["content"]).decode()
    result=[];group=None
    for line in text.splitlines():
        heading=re.match(r"### 组 (\d+)",line)
        if heading: group=heading[1]
        row=re.match(r"\| (\d+) \| `([^`]+)` \| ([^|]+) \| ([^|]+) \|",line)
        if row and group:
            result.append(dict(case_id=f"ST{group}_{int(row[1]):02d}",previous_IT_case=row[2],
                               readiness=row[3].strip(),description=row[4].strip(),source_commit=commit))
    if len(result)!=101: raise ValueError("ST catalog changed; review the source before exporting")
    for number in range(6,10):
        result.append(dict(case_id=f"smoke{number:02d}",previous_IT_case=f"00_{number:02d}_"+
                           {6:"loopback_mmio",7:"loopback_irq",8:"compute_irq",9:"compute_mmio"}[number],
                           readiness="独立Linux smoke已实现；07/08使用MMIO而非PLIC中断",
                           description="N4096回环" if number<8 else "N4096 PMUL",source_commit=commit))
    return result


def output(args,records,by_id):
    from openpyxl import Workbook,load_workbook
    from openpyxl.styles import Alignment,Font,PatternFill
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.utils import get_column_letter
    book=Workbook();book.remove(book.active)
    am_revision=revision(REPO)
    def sheet(name,headers,data,checks=()):
        ws=book.create_sheet(name);ws.append(headers)
        for row in data: ws.append([json.dumps(v,ensure_ascii=False) if isinstance(v,(list,dict)) else v for v in row])
        ws.freeze_panes="C2";ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:
            cell.fill=PatternFill('solid',fgColor='17365D');cell.font=Font(color='FFFFFF',bold=True)
            cell.alignment=Alignment(wrap_text=True,vertical='center')
        ws.row_dimensions[1].height=42
        for col,title in enumerate(headers,1):
            ws.column_dimensions[get_column_letter(col)].width=16 if col in checks else min(62,max(18,len(title)*2))
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment=Alignment(vertical='top',wrap_text=True)
                if cell.row%2==0: cell.fill=PatternFill('solid',fgColor='F1F5FA')
        for col in checks:
            letter=get_column_letter(col);area=f'{letter}2:{letter}{ws.max_row}'
            validation=DataValidation(type='list',formula1='"✓,✗,待测,不适用"',allow_blank=True)
            ws.add_data_validation(validation);validation.add(area)
            ws.conditional_formatting.add(area,CellIsRule(operator='equal',formula=['"✓"'],fill=PatternFill('solid',fgColor='C6EFCE'),font=Font(color='006100')))
            ws.conditional_formatting.add(area,CellIsRule(operator='equal',formula=['"✗"'],fill=PatternFill('solid',fgColor='FFC7CE'),font=Font(color='9C0006')))
        return ws
    definitions=[
      ["用途","参数已据交付填充；IT/ST通过勾选、实测时间和实测AM/RTL版本留空，由测试人员填写。"],
      ["通过判据","模型PASS/编译PASS/可自检不等于IT或ST通过。填写✓后变绿色；故障/负向用例按其预期结果验收。"],
      ["N","多项式次数；每个RNS分量有N个32bit系数。N65536是demo参数，不代表安全强度提升。"],
      ["L","SEAL first data context的chain_index。初始Q数量通常为L+1；不要把L、Q数量、密文分量混用。原语不定义L。"],
      ["Q/P/QP","Q为当前数据模数链，P为KeySwitch特殊基；QP是拼接的模数基，另列精确整数乘积。"],
      ["dnum","真实评估密钥digit向量长度。未使用评估密钥填不适用。legacy Auto为2；SEAL以原生密钥对象为准。"],
      ["RNS分量","一条多项式有|Q|个RNS分量；2分量密文有2×|Q|条多项式。三个算法的Q/P并不相同。"],
      ["golden","05使用inline-asm自带modified-SEAL；逐原始物理word精确比较，不能用解密近似相等掩盖reference不一致。"],
      ["计时区间","07每个ELF一次：首条HPU发令前rdcycle，到末尾PSYNC完成且BUSY清零后rdcycle；统一在自检后打印。"],
      ["计时不包含","镜像拷贝/cache clean、CSR初始化、golden比较、guard检查、UART打印。包含CPU发令/反压、DLOAD/DSTORE、HPU计算及完成轮询；不是纯HPU核心cycle。"],
      ["对比","HPU秒=cycles/目标SCPU频率Hz；Speedup=188服务器CPU秒/HPU秒。不能拿不同频率主机cycle直接相除；未实测不填10x。"],
      ["重复测量","07默认不串行塞多轮。建议同一ELF至少独立运行3次，记录每轮再算平均/最小/最大/波动。"],
      ["内存","新增N65536窗口0x90000000，RAM至少512MiB；VCS使用+ram_size=536870912。软件轮询budget不等于仿真max-cycle预算。"],
      ["安全参数","本交付SEALContext(sec_level_type::none)，security_level_bits=0；仅功能/性能demo，不能用于真实密文安全性声明。"],
      ["ST","System-test仍有未接入producer的占位项。ST页的对应IT编号仅作映射，不能继承IT通过标记或宣称相同执行路径。"],
      ["版本","AM源码版本="+am_revision+"；实测版本另填，不能由当前HEAD反推此前通过版本。"],
      ["CPU复现","cmake目标hpu_cpu_benchmark：SCHEME OP N ROUNDS。固定输入/种子/密钥形状；解码、加解密、密钥生成在计时外；逐轮raw-word验证。"],
      ["旧07","N4096旧性能入口保存在archive/performance_before_n65536，历史编号不映射成新的N65536通过结果。"],
    ]
    ws=sheet('填写说明',['字段','定义与限制'],definitions);ws.column_dimensions['B'].width=110
    inventory=[];parameters=[]
    for record in records:
        p=record['parameters'];f=p.get('fixture',{});q=p.get('Q',[]);special=p.get('P',[])
        inventory.append([record['case_id'],GROUPS[int(record['case_id'][:2])],None,None,record['qualifier'],
            p.get('scheme','不适用'),p.get('N'),p.get('L'),len(q) or None,len(special) or None,p.get('dnum'),
            '; '.join(map(str,q)), '; '.join(map(str,special)),record['source'],None,None,None,None,p.get('degree_source','producer metadata')])
        parameters.append([record['case_id'],p.get('program_stem','源码直接发令'),p.get('scheme'),p.get('operation'),p.get('N'),p.get('L'),
            len(q) or None,len(special) or None,len(q+special) or None,p.get('dnum') or '不适用',q,special,
            [v.bit_length() for v in q],[v.bit_length() for v in special],product(q),product(special),product(q+special),
            p.get('plain_modulus'),f.get('input_components'),f.get('output_components'),p.get('input_limb_counts'),
            p.get('output_Q'),f.get('input_scale'),f.get('right_scale'),f.get('output_scale'),
            f.get('input_correction_factor'),f.get('right_correction_factor'),f.get('output_correction_factor'),
            f.get('rotation_steps'),3 if f.get('rotation_steps') else None,
            f.get('bfv_levels'),p.get('modulus_table'),p.get('evaluation_key_limbs'),f.get('seed'),
            p.get('memory_base'),p.get('used_lines'),p.get('guard_lines'),p.get('image_bytes'),p.get('golden_bytes'),
            p.get('minimum_ram_bytes'),p.get('instruction_count'),p.get('dma_count'),p.get('golden_source'),
            p.get('producer_commit'),p.get('source_parameters',p.get('degree_source')),p.get('operation_graph'),p.get('note'),
            p.get('security_level_bits'),p.get('input_recipe'),p.get('input_domain'),p.get('rotation_galois_element')])
    sheet('IT验收清单',['用例编号','分组','IT通过','ST通过','实现状态（非实测）','方案','N','L(chain_index)','Q-RNS数','P-RNS数','dnum',
          'Q模数','P模数','源码','实测耗时秒','实测AM版本','实测RTL版本','测试人/日期/日志','参数来源或限制'],inventory,(3,4))
    sheet('完整密文及降级参数',['用例编号','指令程序','方案','操作','N','L','Q分量数','P分量数','QP分量数','dnum','Q primes','P primes',
          'Q位宽','P位宽','Q乘积(十进制文本)','P乘积(十进制文本)','QP乘积(十进制文本)','明文t','输入密文分量','输出密文分量',
          '各输入RNS多项式数','各阶段输出Q数','输入scale','右输入scale','输出scale','输入CF','右输入CF','输出CF','rotation steps','generator',
          'BFV各层B及m_sk','完整MOD_ID模表','评估密钥RNS多项式数','确定性种子','窗口基址','窗口line数','guard line数','初始镜像bytes','golden bytes',
          '最小RAM bytes','HPU指令数','DMA次数','golden来源','inline-asm commit','参数文件','SEAL操作图','限制',
          '安全等级bits（0=demo）','明文输入构造','输入物理域','Galois element'],parameters)
    children=[]
    for row in table(ROOT/'subtests/cases.tsv'):
        children.append([row['subtest_id'],row['parent_case_id'],None,None,int(row['subcase']),int(row['programs']),row['description'],4096,'见父用例参数/该subcase源码',row['source']])
    sheet('03独立子用例',['用例编号','父用例','IT通过','ST通过','mainargs subcase','程序段数','检查点','N','参数说明','源码'],children,(3,4))
    perf=[]
    for row in table(ROOT/'perf-cases.tsv'):
        p=by_id[row['case_id']]['parameters']
        for run in range(1,4):
            perf.append([row['case_id'],row['functional_case_id'],run,None,65536,p.get('scheme'),len(p['Q']),len(p['P']),p.get('dnum'),
                         None,None,None,None,None,None,None,None,None,None])
    ws=sheet('07性能逐轮记录',['用例编号','对应功能用例','独立轮次','功能通过','N','方案','Q-RNS数','P-RNS数','dnum',
             'cycle_start','cycle_end','cycle_delta','SCPU频率Hz','HPU秒','CPU188秒','Speedup','AM/RTL版本','CPU型号/编译参数','日志位置'],perf,(4,))
    for i in range(2,ws.max_row+1):
        ws.cell(i,12,f'=IF(OR(J{i}="",K{i}=""),"",K{i}-J{i})')
        ws.cell(i,14,f'=IF(OR(L{i}="",M{i}="",M{i}<=0),"",L{i}/M{i})')
        ws.cell(i,16,f'=IF(OR(N{i}="",O{i}="",N{i}<=0),"",O{i}/N{i})')
    if args.st_github or args.st_snapshot:
        st=load(args.st_snapshot) if args.st_snapshot else fetch_st(args.gh)
        aliases={r['previous_case_id']:r['case_id'] for r in table(ROOT/'case-aliases.tsv')}
        st_rows=[]
        for row in st:
            target=aliases.get(row['previous_IT_case'],row['previous_IT_case'])
            p=by_id.get(target,{}).get('parameters',{})
            st_rows.append([row['case_id'],target,None,row['readiness'],row['description'],p.get('N'),
                '计划参数，producer尚未接入ST' if 'producer' in row['readiness'] else 'ST自身数据/同步路径，以ST源码为准',
                None,None,None,None,f"https://github.com/cuihu2/System-test/blob/{row['source_commit']}/CASE_STATUS.md"])
        sheet('ST验收清单',['ST编号','对应IT编号（非通过继承）','ST通过','源码实现状态','检查点','N或计划N','参数状态','实测秒','ST版本','RTL版本','测试人/日志','来源'],st_rows,(3,))
        (args.output/'ST_INVENTORY.json').write_text(json.dumps(st,ensure_ascii=False,indent=2)+'\n')
    archived=table(ROOT/'archived-cases.tsv')
    sheet('旧07归档',['历史编号','实现状态','归档源码','说明'],[[r['case_id'],r['qualifier'],r['source'],'不映射成N65536；通过结果不自动继承'] for r in archived])
    filename=args.output/'HPU_IT_ST用例与参数清单.xlsx';book.save(filename)
    # 文件结构/空验收格/验证规则检查；没有桌面Excel时不声称已做桌面渲染测试。
    check=load_workbook(filename)
    assert check['IT验收清单'].max_row==len(records)+1
    assert all(check['IT验收清单'].cell(i,c).value is None for i in range(2,len(records)+2) for c in (3,4))
    assert len(check['IT验收清单'].data_validations.dataValidation)==2
    (args.output/'CASE_PARAMETERS.json').write_text(json.dumps(dict(am_source_revision=am_revision,cases=records),ensure_ascii=False,indent=2)+'\n')
    with (args.output/'IT验收清单.csv').open('w',encoding='utf-8-sig',newline='') as stream:
        writer=csv.writer(stream);writer.writerow([c.value for c in check['IT验收清单'][1]]);writer.writerows(inventory)
    print(f'Inventory: {len(records)} IT entries, {len(children)} subtests; pass cells blank; {filename}')


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    for key in ('generated','legacy','parameter-tool','output'): parser.add_argument('--'+key,type=Path,required=True)
    parser.add_argument('--st-github',action='store_true',help='仅用已授权GitHub账号读取System-test清单')
    parser.add_argument('--st-snapshot',type=Path,help='复用带commit的ST_INVENTORY.json，不反复读取未变化源')
    parser.add_argument('--gh',default='gh')
    args=parser.parse_args();args.output.mkdir(parents=True,exist_ok=True)
    records,by_id=collect(args);output(args,records,by_id)
