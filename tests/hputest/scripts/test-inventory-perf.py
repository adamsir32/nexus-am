#!/usr/bin/env python3
"""用实际交付包核验07五种规模、原level/L16的参数和Excel逐轮记录。"""
import argparse
from collections import Counter
import importlib.util
from pathlib import Path
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import patch

HERE=Path(__file__).resolve().parent
spec=importlib.util.spec_from_file_location('inventory',HERE/'export-case-inventory.py')
inventory=importlib.util.module_from_spec(spec);spec.loader.exec_module(inventory)


class PerformanceInventory(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temporary=tempfile.TemporaryDirectory(prefix='hpu-inventory-perf-',dir=ARGS.generated.parent)
        cls.args=SimpleNamespace(generated=ARGS.generated,output=Path(cls.temporary.name),
                                 parameter_tool=ARGS.parameter_tool,
                                 legacy=Path('/unused/legacy'),st_github=False,st_snapshot=None)
        cls.perf=inventory.table(inventory.ROOT/'perf-cases.tsv')
        # 不加载对应功能用例：即使functional_case_id为空，也必须独立读取交付参数。
        read_cases=inventory.select_cases
        selected={row['case_id'] for row in cls.perf}
        selected.update(('03_16_pmodld_n65536','05_07_ckks_hmul_reline_n4096',
                         '05_41_ckks_hmul_reline_n65536'))
        def subset(scope):
            return [row for row in read_cases(scope) if row['case_id'] in selected]
        with patch.object(inventory,'select_cases',side_effect=subset):
            cls.records,cls.by_id=inventory.collect(cls.args)

    @classmethod
    def tearDownClass(cls):
        cls.temporary.cleanup()

    def test_performance_parameters_match_each_delivery(self):
        self.assertEqual(len(self.perf),200)
        self.assertEqual(Counter(int(row['degree']) for row in self.perf),
                         {4096:40,8192:40,16384:40,32768:40,65536:40})
        self.assertEqual(Counter('_l16_' in row['program_stem'] for row in self.perf),
                         {False:110,True:90})
        self.assertEqual(Counter((int(row['degree']),'_l16_' in row['program_stem']) for row in self.perf),
                         {(degree,is_l16):18 if is_l16 else 22
                          for degree in (4096,8192,16384,32768,65536) for is_l16 in (False,True)})
        self.assertEqual(sum(not row['functional_case_id'] for row in self.perf),178)
        for row in self.perf:
            with self.subTest(case=row['case_id']):
                p=self.by_id[row['case_id']]['parameters']
                self.assertEqual(p['N'],int(row['degree']))
                self.assertEqual(p['program_stem'],row['program_stem'])
                self.assertEqual(p['metric'],'issue_to_done_v2')
                self.assertEqual(len(p['Q']),int(row['level']))
                self.assertEqual(int(p['memory_base'],0),int(row['memory_base'],0))
                self.assertEqual(p['minimum_ram_bytes'],int(row['minimum_ram_bytes']))
                self.assertLessEqual(int(p['memory_base'],0)+p['image_bytes'],
                                     0x80000000+p['minimum_ram_bytes'])
                if '_l16_' in row['program_stem']:
                    self.assertEqual(len(p['Q']),16)
                    self.assertEqual(len(p['P']),1)
                    self.assertEqual(p['fixture']['Q_rns_count'],16)
                    self.assertEqual(p['memory_base'],'0xc0000000')
                    self.assertEqual(p['minimum_ram_bytes'],2147483648)
                else:
                    self.assertEqual(p['memory_base'],'0x90000000')
                    self.assertEqual(p['minimum_ram_bytes'],536870912)
                self.assertEqual(p['functional_case_id'],row['functional_case_id'] or None)
                directory=ARGS.generated/row['delivery']/row['program_stem']
                if row['delivery']=='large-data':
                    original=inventory.load(directory/'PARAMETERS.json')
                    if original['operation'] in ('ntt','intt'):
                        self.assertEqual(p['Q'],original['q_moduli'][:1])
                        self.assertEqual(p['modulus_table'],original['q_moduli'])
                    else:
                        self.assertEqual(p['Q'],original['q_moduli'])
                    self.assertEqual(p['P'],original['p_moduli'])
                else:
                    upstream=directory/'upstream'
                    package=inventory.load(upstream/'package.json')
                    original=inventory.load(upstream/package['parameters'])
                    self.assertEqual(p['N'],original['poly_modulus_degree'])
                    self.assertEqual(p['Q'],original['levels'][0]['moduli'])
                    self.assertEqual(p['Q']+p['P'],original['key_moduli'])
                self.assertIsNone(self.by_id[row['case_id']]['IT_pass'])
                self.assertIsNone(self.by_id[row['case_id']]['ST_pass'])

    def test_functional_parameters_keep_their_semantics(self):
        pmod=self.by_id['03_16_pmodld_n65536']['parameters']
        self.assertEqual(len(pmod['Q']),1)
        self.assertEqual(pmod['modulus_table'],pmod['q_moduli'])
        functional=self.by_id['05_41_ckks_hmul_reline_n65536']['parameters']
        perf_id=next(row['case_id'] for row in self.perf
                     if row['program_stem']=='seal_ckks_hmul_n65536')
        timed=self.by_id[perf_id]['parameters']
        self.assertEqual(functional,{key:value for key,value in timed.items()
                                    if key not in ('functional_case_id','metric')})
        self.assertNotIn('metric',pmod)
        # N4096共享数据包，但05沿用原窗口，07使用性能窗口，参数不能串改。
        functional=self.by_id['05_07_ckks_hmul_reline_n4096']['parameters']
        perf_id=next(row['case_id'] for row in self.perf
                     if row['program_stem']=='seal_ckks_hmul_n4096')
        timed=self.by_id[perf_id]['parameters']
        self.assertEqual(functional['memory_base'],'0x87000000')
        self.assertEqual(functional['minimum_ram_bytes'],268435456)
        self.assertEqual(timed['memory_base'],'0x90000000')
        self.assertEqual(timed['minimum_ram_bytes'],536870912)
        self.assertEqual({key:value for key,value in functional.items()
                          if key not in ('memory_base','minimum_ram_bytes')},
                         {key:value for key,value in timed.items()
                          if key not in ('functional_case_id','metric','memory_base','minimum_ram_bytes')})

    def test_l16_native_parameter_fallback_keeps_requested_q_count(self):
        stem='seal_ckks_hadd_l16_n4096'
        producer=inventory.application_specs()[stem]
        fixture=ARGS.generated/'application-data'/stem/'FIXTURE_PARAMETERS.json'
        cache=self.args.output/'fallback-l16';cache.mkdir()
        exists=Path.exists
        # 只屏蔽这一个元数据文件，模拟旧交付缺少缓存；不修改交付目录。
        with patch.object(Path,'exists',lambda path: False if path==fixture else exists(path)), \
             patch.object(inventory.subprocess,'check_output',wraps=inventory.subprocess.check_output) as native:
            params=inventory.application(producer,ARGS.generated,ARGS.parameter_tool,cache)
        native.assert_called_once_with([str(ARGS.parameter_tool),'ckks','hadd','4096','library','16'])
        self.assertEqual(params['fixture']['Q_rns_count'],16)
        with self.assertRaisesRegex(ValueError,'Q RNS count'):
            inventory.application(dict(producer,q_rns_count=4),ARGS.generated,ARGS.parameter_tool,cache)

    def test_wrong_level_or_too_small_ram_contract_is_rejected(self):
        case=next(row for row in self.perf if row['program_stem']=='seal_ckks_hadd_n4096')
        case_source=next(row for row in inventory.select_cases('all') if row['case_id']==case['case_id'])
        read_table=inventory.table
        for field,value,error in (('level','1','Q RNS count'),
                                  ('minimum_ram_bytes','268435456','RAM contract')):
            with self.subTest(field=field):
                def changed(path,delimiter='\t'):
                    rows=read_table(path,delimiter)
                    if path.name=='perf-cases.tsv':
                        return [dict(case,**{field:value})]
                    return rows
                with patch.object(inventory,'select_cases',return_value=[case_source]), \
                     patch.object(inventory,'table',side_effect=changed), \
                     self.assertRaisesRegex(ValueError,error):
                    inventory.collect(self.args)

    def test_excel_rounds_use_actual_degree_and_leave_measurements_blank(self):
        from openpyxl import load_workbook
        inventory.output(self.args,self.records,self.by_id)
        book=load_workbook(self.args.output/'HPU_IT_ST用例与参数清单.xlsx')
        sheet=book['07性能逐轮记录']
        rows=list(sheet.iter_rows(min_row=2,values_only=True))
        self.assertEqual(len(rows),600)
        by_case={row['case_id']:row for row in self.perf}
        for row in rows:
            source=by_case[row[0]]
            self.assertEqual(row[1],source['functional_case_id'] or None)
            self.assertEqual(row[4],int(source['degree']))
            self.assertEqual(row[6],len(self.by_id[row[0]]['parameters']['Q']))
            self.assertEqual(row[7],len(self.by_id[row[0]]['parameters']['P']))
            self.assertIn(row[2],(1,2,3))
            for column in (3,9,10,12,14):
                self.assertIsNone(row[column])
        self.assertEqual(Counter(row[4] for row in rows),
                         {4096:120,8192:120,16384:120,32768:120,65536:120})


class InventoryScopes(unittest.TestCase):
    def setUp(self):
        self.temporary=tempfile.TemporaryDirectory(prefix='hpu-inventory-scope-',dir=ARGS.generated.parent)
        self.addCleanup(self.temporary.cleanup)

    def collect_and_export(self,scope):
        from openpyxl import load_workbook
        selected=inventory.select_cases(scope)
        selected_ids={row['case_id'] for row in selected}
        allowed_applications=set(inventory.application_specs(scope))
        allowed_large={row['program_stem'] for row in inventory.table(inventory.ROOT/'large-cases.tsv')
                       if row['case_id'] in selected_ids}
        allowed_large.update(row['program_stem'] for row in inventory.table(inventory.ROOT/'perf-cases.tsv')
                             if row['case_id'] in selected_ids and row['delivery']=='large-data')
        application=inventory.application;large=inventory.large_parameters
        def read_application(spec,*args):
            self.assertIn(spec['program_stem'],allowed_applications)
            return application(spec,*args)
        def read_large(stem,*args):
            self.assertIn(stem,allowed_large)
            return large(stem,*args)
        output=Path(self.temporary.name)/scope.replace(':','_');output.mkdir()
        args=SimpleNamespace(generated=ARGS.generated,output=output,
                             parameter_tool=ARGS.parameter_tool,
                             legacy=ARGS.legacy if scope=='functional' else Path('/no-functional-data'),
                             scope=scope,st_github=False,st_snapshot=None)
        with patch.object(inventory,'application',side_effect=read_application), \
             patch.object(inventory,'large_parameters',side_effect=read_large):
            records,by_id=inventory.collect(args)
        self.assertEqual(set(by_id),selected_ids)
        inventory.output(args,records,by_id)
        book=load_workbook(output/'HPU_IT_ST用例与参数清单.xlsx')
        self.assertEqual({row[0] for row in book['IT验收清单'].iter_rows(min_row=2,values_only=True)},selected_ids)
        expected_perf={case_id for case_id in selected_ids if case_id.startswith('07_')}
        performance=list(book['07性能逐轮记录'].iter_rows(min_row=2,values_only=True))
        self.assertEqual(Counter(row[0] for row in performance),{case_id:3 for case_id in expected_perf})
        self.assertEqual({row['case_id'] for row in inventory.load(output/'CASE_PARAMETERS.json')['cases']},selected_ids)
        return selected_ids

    def test_functional_export_does_not_read_performance_only_data(self):
        if ARGS.legacy is None:
            self.skipTest('functional scope requires --legacy with N4096 producer parameters')
        selected=self.collect_and_export('functional')
        self.assertTrue(selected)
        self.assertEqual(selected,{row['case_id'] for row in inventory.table(inventory.ROOT/'cases.tsv')
                                   if not row['case_id'].startswith('07_')})

    def test_representative_performance_shards_export_only_selected_cases(self):
        for scope,count in (('performance:4096:baseline',22),('performance:8192:l16-ckks',6)):
            with self.subTest(scope=scope):
                self.assertEqual(len(self.collect_and_export(scope)),count)
        first=self.collect_and_export('performance:65536:l16-bfv-a')
        second=self.collect_and_export('performance:65536:l16-bfv-b')
        self.assertEqual(len(first),3);self.assertEqual(len(second),3)
        self.assertFalse(first&second)
        self.assertEqual(first|second,{row['case_id'] for row in inventory.select_cases('performance:65536:l16-bfv')})


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--generated',type=Path,required=True)
    parser.add_argument('--parameter-tool',type=Path,required=True)
    parser.add_argument('--legacy',type=Path,help='functional导出验证用的N4096生产者参数根目录')
    ARGS,extra=parser.parse_known_args()
    unittest.main(argv=[__file__,*extra])
